"""Uploaded files the interview is based on (CV, job description, notes, course material...)."""
from __future__ import annotations

import json
import re
import threading
import time
import uuid
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".html", ".htm"}
MAX_TOTAL_CHARS = 400_000  # roughly 100k tokens; more than this is trimmed with a warning


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            # Many PDFs come out as "word\n \nword": every word on its own line, 2x the tokens on
            # every answer. Drop that padding only; layout mode would lose some text (forms, rotated).
            text = re.sub(r"\n[ \t]+\n", " ", page.extract_text() or "")
            lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
            text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
            if text:
                pages.append(f"[page {i}]\n{text}")
        return "\n\n".join(pages)
    if suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if suffix in (".html", ".htm"):
        raw = path.read_text(encoding="utf-8", errors="replace")
        raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
        return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", raw)).strip()
    if suffix in SUPPORTED:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def _safe_name(name: str) -> str:
    name = Path(name).name
    return re.sub(r"[^\w.\- ()]+", "_", name)[:120] or "file"


class DocumentStore:
    def __init__(self, folder: Path):
        self.folder = folder
        self.folder.mkdir(parents=True, exist_ok=True)
        self.index_path = folder / "index.json"
        self._lock = threading.Lock()
        self._docs: list[dict] = []
        if self.index_path.exists():
            try:
                self._docs = json.loads(self.index_path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                self._docs = []
        self._docs = [d for d in self._docs if (self.folder / f"{d['id']}.txt").exists()]

    def _save(self) -> None:
        self.index_path.write_text(json.dumps(self._docs, indent=2), encoding="utf-8")

    def list(self) -> list[dict]:
        with self._lock:
            return [dict(d) for d in self._docs]

    def add_file(self, filename: str, data: bytes) -> dict:
        name = _safe_name(filename)
        suffix = Path(name).suffix.lower()
        if suffix not in SUPPORTED:
            raise ValueError(f"{name}: only PDF, Word (.docx), text, Markdown, CSV, JSON or HTML files are supported.")
        doc_id = uuid.uuid4().hex[:10]
        original = self.folder / f"{doc_id}{suffix}"
        original.write_bytes(data)
        try:
            text = extract_text(original).strip()
        except Exception as e:
            original.unlink(missing_ok=True)
            raise ValueError(f"{name}: couldn't read text from this file ({e}).") from e
        if not text:
            original.unlink(missing_ok=True)
            raise ValueError(f"{name}: no readable text found (a scanned PDF needs OCR first).")
        return self._register(doc_id, name, text)

    def add_text(self, name: str, text: str) -> dict:
        text = (text or "").strip()
        if not text:
            raise ValueError("The pasted text is empty.")
        return self._register(uuid.uuid4().hex[:10], _safe_name(name or "Pasted notes") , text)

    def _register(self, doc_id: str, name: str, text: str) -> dict:
        (self.folder / f"{doc_id}.txt").write_text(text, encoding="utf-8")
        doc = {"id": doc_id, "name": name, "chars": len(text), "words": len(text.split()), "added": time.time()}
        with self._lock:
            self._docs.append(doc)
            self._save()
        return doc

    def remove(self, doc_id: str) -> bool:
        with self._lock:
            before = len(self._docs)
            self._docs = [d for d in self._docs if d["id"] != doc_id]
            self._save()
        for f in self.folder.glob(f"{doc_id}.*"):
            f.unlink(missing_ok=True)
        return len(self._docs) < before

    def combined(self) -> tuple[str, bool]:
        """All documents as one block for the AI prompt. Returns (text, was_trimmed)."""
        parts, total, trimmed = [], 0, False
        for d in self.list():
            try:
                text = (self.folder / f"{d['id']}.txt").read_text(encoding="utf-8")
            except OSError:
                continue
            room = MAX_TOTAL_CHARS - total
            if room <= 0:
                trimmed = True
                break
            if len(text) > room:
                text, trimmed = text[:room] + "\n[...trimmed...]", True
            parts.append(f'<document name="{d["name"]}">\n{text}\n</document>')
            total += len(text)
        return "\n\n".join(parts), trimmed
